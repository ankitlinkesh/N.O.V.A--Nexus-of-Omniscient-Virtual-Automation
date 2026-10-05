"""Standalone verifier for Phase 132 (NOVA reads real PDF and Word files).

Builds REAL documents in a temp home (no mocked extractor) and reads them through
the registry's file.read_text:
1. pypdf and python-docx are installed and pinned in requirements.txt.
2. A PDF with a text layer is read, and the result is still untrusted content.
3. A scanned PDF (no text layer) and a password-protected PDF are refused with the
   reason, never returned as an "empty" read; the protected text never leaks.
4. A PDF restricted only by an owner password is read.
5. A .docx is read with its tables, in document order (tables were skipped).
6. README records the phase.
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend" / "tests"))
failures = 0


def emit(case: str, ok: bool, **extra: object) -> int:
    payload = {"case": case, "pass": bool(ok)}
    payload.update(extra)
    print(json.dumps(payload, indent=2, default=str))
    return 0 if ok else 1


try:
    import docx
    import pypdf

    from backend.eva.tools.registry import ToolRegistry
    from test_phase132_read_documents import _text_pdf

    requirements = (ROOT / "requirements.txt").read_text(encoding="utf-8")
    failures += emit("pypdf and python-docx installed and pinned", "pypdf==" in requirements and "python-docx==" in requirements, pypdf=pypdf.__version__)

    home = Path(tempfile.mkdtemp(prefix="nova_p132_"))
    downloads = home / "Downloads"
    downloads.mkdir()
    real_home = Path.home
    Path.home = classmethod(lambda cls: home)  # type: ignore[method-assign]
    try:
        def read(name: str) -> dict:
            return ToolRegistry().run("file.read_text", path=f"Downloads/{name}")

        (downloads / "invoice.pdf").write_bytes(_text_pdf("Invoice total 4200 rupees"))
        r = read("invoice.pdf")
        failures += emit("a PDF with a text layer is read, as untrusted content", r.get("ok") is True and "Invoice total 4200 rupees" in r.get("text", "") and r.get("untrusted") is True)

        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with (downloads / "scan.pdf").open("wb") as handle:
            writer.write(handle)
        r = read("scan.pdf")
        failures += emit("a scanned PDF is refused with the reason", r.get("ok") is False and r.get("error") == "no_text" and "scanned" in r.get("message", ""), message=r.get("message"))

        (downloads / "plain.pdf").write_bytes(_text_pdf("secret plan"))
        for name, user_pw in (("locked.pdf", "hunter2"), ("restricted.pdf", "")):
            w = pypdf.PdfWriter()
            for page in pypdf.PdfReader(str(downloads / "plain.pdf")).pages:
                w.add_page(page)
            w.encrypt(user_password=user_pw, owner_password="owner", algorithm="AES-128")
            with (downloads / name).open("wb") as handle:
                w.write(handle)
        locked, restricted = read("locked.pdf"), read("restricted.pdf")
        failures += emit("a password-protected PDF is refused and its text never leaks", locked.get("ok") is False and "password" in locked.get("message", "") and "secret plan" not in json.dumps(locked))
        failures += emit("an owner-password-only PDF is still read", restricted.get("ok") is True and "secret plan" in restricted.get("text", ""))

        document = docx.Document()
        document.add_paragraph("Quarterly budget")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Cost"
        table.cell(1, 0).text, table.cell(1, 1).text = "Laptop", "55000"
        document.add_paragraph("Approved by finance")
        document.save(str(downloads / "budget.docx"))
        r = read("budget.docx")
        lines = [line for line in r.get("text", "").split("\n") if line.strip()]
        failures += emit("a .docx is read with its tables in document order", lines == ["Quarterly budget", "Item | Cost", "Laptop | 55000", "Approved by finance"], lines=lines)
    finally:
        Path.home = real_home  # type: ignore[method-assign]

    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    failures += emit("README records Phase 132", "| 132 |" in readme)
except Exception as exc:  # a crash is a failure, never a pass
    failures += emit("verifier crashed", False, error=f"{type(exc).__name__}: {exc}")

print(json.dumps({"overall_pass": failures == 0, "failures": failures}, indent=2))
sys.exit(0 if failures == 0 else 1)
